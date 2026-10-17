"""Generate the synthetic ingestion corpus in this directory.

Every organisation, person, place and number below is invented. The files exist
to exercise ingestion (P0-09, P0-10) and parsing spot checks (P0-11) end to end
without real documents. Regenerate with:

    uv run --no-project --with reportlab==5.0.1 --with python-docx==1.2.0 \
        --with python-pptx==1.0.2 python packages/eval/fixtures/corpus/make_corpus.py

Output is not byte-for-byte reproducible (the formats embed timestamps); the
committed files are the reference copies.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).resolve().parent


def make_pdf() -> None:
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Tidewater Offshore Wind Outlook 2026", styles["Title"]),
        Paragraph("Fictional Institute for Coastal Energy (FICE), synthetic report", styles["Normal"]),
        Spacer(1, 12),
        Paragraph("1. Summary", styles["Heading2"]),
        Paragraph(
            "Installed offshore wind capacity in the fictional Tidewater region reached "
            "4.2 GW at the end of 2025. FICE projects 9.8 GW by 2030 if the Marlow Bank "
            "auction clears on schedule. Grid connection queues remain the main constraint; "
            "the average wait for a connection offer rose from 14 to 22 months.",
            styles["BodyText"],
        ),
        Paragraph("2. Capacity by zone", styles["Heading2"]),
        Paragraph("Table 1 lists installed and planned capacity by zone.", styles["BodyText"]),
        Spacer(1, 6),
    ]
    table = Table(
        [
            ["Zone", "Installed 2025 (GW)", "Planned 2030 (GW)"],
            ["Marlow Bank", "1.1", "3.9"],
            ["Grey Shoal", "2.0", "3.4"],
            ["Esk Deep", "1.1", "2.5"],
        ]
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ]
        )
    )
    story += [
        table,
        PageBreak(),
        Paragraph("3. Costs", styles["Heading2"]),
        Paragraph(
            "The levelised cost of electricity for fixed-bottom projects fell to 61 EUR/MWh "
            "in the synthetic 2025 auction round, down from 74 EUR/MWh in 2023. Floating "
            "projects remain above 120 EUR/MWh. FICE attributes the decline to larger "
            "turbines (15 MW class) and lower financing costs.",
            styles["BodyText"],
        ),
        Paragraph("4. Risks", styles["Heading2"]),
        Paragraph(
            "Supply-chain bottlenecks for export cables and installation vessels are the "
            "largest schedule risk. A two-year slip at Marlow Bank would cut the 2030 "
            "projection to 7.6 GW.",
            styles["BodyText"],
        ),
    ]
    SimpleDocTemplate(str(OUT / "tidewater-offshore-wind-2026.pdf"), pagesize=A4).build(story)


def make_docx() -> None:
    doc = Document()
    doc.add_heading("Analysis note: heat pump adoption in Brightwater", 0)
    doc.add_paragraph("Synthetic internal analysis. All figures are invented.")
    doc.add_heading("Key findings", level=1)
    doc.add_paragraph(
        "Heat pump installations in the fictional municipality of Brightwater doubled "
        "between 2023 and 2025, from 3,100 to 6,300 units per year."
    )
    doc.add_paragraph(
        "The subsidy scheme covering 40 percent of installation costs explains most of "
        "the increase; installer capacity is now the binding constraint."
    )
    doc.add_heading("Recommendation", level=1)
    doc.add_paragraph(
        "Prioritise installer training over further subsidy increases; each additional "
        "certified installer adds roughly 85 installations per year."
    )
    doc.save(str(OUT / "brightwater-heat-pumps-note.docx"))


def make_pptx() -> None:
    prs = Presentation()
    slides = [
        ("Battery storage in Kestrel Valley", "Synthetic briefing deck", None),
        (
            "Deployed capacity",
            "Grid batteries in the fictional Kestrel Valley reached 850 MWh in 2025.",
            "Speaker note: capacity tripled since 2023.",
        ),
        (
            "Revenue stack",
            "Frequency response provides 55 percent of revenue; arbitrage 30 percent.",
            "Speaker note: arbitrage share is rising.",
        ),
        (
            "Outlook",
            "Kestrel Valley Energy Board expects 2.4 GWh of storage by 2029.",
            None,
        ),
    ]
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    box = prs.slides[3].shapes.add_textbox(Inches(1), Inches(6), Inches(8), Inches(1))
    box.text_frame.text = "Source: Kestrel Valley Energy Board (fictional)"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    prs.save(str(OUT / "kestrel-valley-storage-briefing.pptx"))


if __name__ == "__main__":
    make_pdf()
    make_docx()
    make_pptx()
